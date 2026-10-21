#!/usr/bin/env python3
"""HUDDLE: organization-scoped cards, relationships and native Android API."""
import security, surfaces, database, object_storage, billing, platform_admin
import argparse,base64,csv,hashlib,hmac,html,io,json,mimetypes,os,re,secrets,sqlite3,sys,threading,time,urllib.error,urllib.parse,urllib.request,zipfile
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
ROOT=Path(__file__).resolve().parent
DATA=Path(os.environ.get('HUDDLE_DATA_DIR',os.environ.get('HELLO_DATA_DIR',ROOT/'data')))
DB=DATA/'huddle.sqlite3'
FIELDS=['name','title','email','phone','bio','photo','personal_links','banners']
LAYOUTS=['classic','portrait','minimal','bold','executive','mint','sunset']
DEMO=False
RATE={};RATE_LOCK=threading.Lock()
class Problem(Exception):
    def __init__(self,message,code=400):self.message,self.code=message,code

def uid():return secrets.token_hex(12)
def now():return int(time.time())
def dump(x):return json.dumps(x,ensure_ascii=False)
def text(v,limit=300):return str(v or '').strip()[:limit]
def connection():
    return database.connection(DB)

def password_hash(v):
    if not isinstance(v,str) or not 10<=len(v)<=256:raise Problem('Use a password between 10 and 256 characters.')
    salt=secrets.token_hex(16);return salt+':'+hashlib.pbkdf2_hmac('sha256',v.encode(),salt.encode(),310000).hex()
def password_matches(v,stored):
    try:
        salt,digest=stored.split(':');return hmac.compare_digest(hashlib.pbkdf2_hmac('sha256',v.encode(),salt.encode(),310000).hex(),digest)
    except (ValueError,AttributeError):return False

def email(v):
    v=text(v,255).lower()
    if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',v) or len(v)>254:raise Problem('Enter a valid email address.')
    return v

def colour(v):
    if not re.fullmatch(r'#[a-fA-F0-9]{6}',str(v)):raise Problem('Use a six-digit hex colour.')
    return v

def social_links(values):
    if not isinstance(values,list) or len(values)>10:raise Problem('Use up to ten links.')
    out=[]
    for x in values:
        label=text(x.get('label'),40);url=text(x.get('url'),1000);p=urllib.parse.urlsplit(url)
        if not label and not url:continue
        if not label or p.scheme not in ('https','http','tel','mailto') or (p.scheme in ('https','http') and not p.netloc):raise Problem('Use a label and a valid contact URL.')
        out.append({'label':label,'url':url})
    return out

def base_url():return os.environ.get('PUBLIC_BASE_URL','http://127.0.0.1:8765').rstrip('/')
def company(c,tenant):
    r=c.execute("SELECT data FROM organizations WHERE id=? AND status='active'",(tenant,)).fetchone()
    if not r:raise Problem('Organization unavailable.',404)
    return json.loads(r[0])
def require_admin(actor):
    if actor['role'] not in ('owner','admin'):raise Problem('Administrator access required.',403)
def require_owner(actor):
    if actor['role']!='owner':raise Problem('Owner access required.',403)
def scoped(c,table,tenant,record_id):
    if table not in ('users','cards','divisions','templates','invitations','requests','nfc_cards','contact_exchanges'):raise ValueError('Invalid repository')
    r=c.execute('SELECT * FROM '+table+' WHERE tenant_id=? AND id=?',(tenant,record_id)).fetchone()
    if not r:raise Problem('Record not found in this organization.',404)
    return r

def audit(c,actor,action,subject):c.execute('INSERT INTO audit(tenant_id,actor,action,subject,created) VALUES(?,?,?,?,?)',(actor['tenant_id'],actor['id'],action,text(subject),now()))
def docs(c,table,tenant):
    if table not in ('templates','divisions'):raise ValueError()
    return [dict(r,data=json.loads(r['data'])) for r in c.execute('SELECT * FROM '+table+' WHERE tenant_id=? ORDER BY created',(tenant,))]

def asset(c,tenant,value):
    value=text(value,220)
    if value and not c.execute('SELECT 1 FROM assets WHERE tenant_id=? AND storage_key=?',(tenant,value)).fetchone():raise Problem('Select an image uploaded to this organization.')
    return value

def publish_assets(c,tenant,values):
    for value in values:
        if value:c.execute("UPDATE assets SET visibility='public' WHERE tenant_id=? AND storage_key=?",(tenant,value))

def user_dict(c,row):
    d=json.loads(row['data']);account=c.execute('SELECT email FROM accounts WHERE id=?',(row['account_id'],)).fetchone()
    d.update({k:row[k] for k in ('id','tenant_id','role','active','created','division_id')});d['email']=account['email']
    card=c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? AND is_default=1',(row['tenant_id'],row['id'])).fetchone()
    if card:
        d.update(json.loads(card['data']));d.update(id=row['id'],tenant_id=row['tenant_id'],role=row['role'],active=row['active'],division_id=row['division_id'],slug=card['slug'],card_id=card['id'],template_id=card['template_id'])
    return d

def create_card(c,tenant,user_id,data,persona='Executive',default=False,slug=None):
    billing.capacity(sys.modules[__name__],c,tenant,'cards')
    member=scoped(c,'users',tenant,user_id);brand=company(c,tenant);tid=data.get('template_id') or brand['default_template'];t=scoped(c,'templates',tenant,tid)
    if not json.loads(t['data']).get('enabled',True):raise Problem('Choose an enabled template.')
    profile={k:data.get(k,[] if k in ('banners','personal_links') else '') for k in FIELDS};profile['banner_mode']=data.get('banner_mode','company')
    cid=uid();slug=slug or re.sub(r'[^a-z0-9]+','-',text(profile['name']).lower()).strip('-')[:40]+'-'+secrets.token_hex(5)
    if default:c.execute('UPDATE cards SET is_default=0 WHERE tenant_id=? AND user_id=?',(tenant,user_id))
    c.execute('INSERT INTO cards VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(cid,tenant,user_id,tid,text(persona,60) or 'Executive',slug,dump(profile),'{}',int(default),1,now(),now()));return cid

def add_user(c,tenant,data,password=None,role='member',account_id=None):
    if role not in ('owner','admin','member'):raise Problem('Choose Owner, Admin or Member.')
    address_for_quota=data.get('email') or (c.execute('SELECT email FROM accounts WHERE id=?',(account_id,)).fetchone()[0] if account_id else '')
    billing.capacity(sys.modules[__name__],c,tenant,'seats',exclude_invitation_email=address_for_quota)
    if role in ('owner','admin'):billing.capacity(sys.modules[__name__],c,tenant,'admins',exclude_invitation_email=address_for_quota)
    if data.get('division_id'):scoped(c,'divisions',tenant,data['division_id'])
    if not account_id:
        address=email(data.get('email'));existing=c.execute('SELECT id FROM accounts WHERE email=?',(address,)).fetchone()
        if existing:raise Problem('An account already exists. Sign in to accept this invitation.',409)
        account_id=uid();c.execute('INSERT INTO accounts VALUES(?,?,?,?,?)',(account_id,address,password_hash(password),now(),None))
    address=c.execute('SELECT email FROM accounts WHERE id=?',(account_id,)).fetchone()['email'];mid=uid()
    profile={'name':text(data.get('name')) or address.split('@')[0],'title':text(data.get('title')),'phone':text(data.get('phone')),'bio':'','photo':'','personal_links':[],'banners':[],'banner_mode':'company'}
    c.execute('INSERT INTO users VALUES(?,?,?,?,?,?,?,?,?)',(mid,tenant,account_id,role,data.get('division_id') or None,dump(profile),1,now(),None))
    c.execute("UPDATE invitations SET status='accepted' WHERE tenant_id=? AND email=? AND status='pending' AND expires>?",(tenant,address,now()))
    create_card(c,tenant,mid,{**profile,**data,'name':profile['name'],'email':address},default=True);return mid

def create_org(c,name,account_id,person_name='',demo_grant=False):
    name=text(name)
    if not name:raise Problem('Enter an organization name.')
    billing.organization_capacity(sys.modules[__name__],c,account_id,enforce=True)
    tenant=uid();slug=re.sub(r'[^a-z0-9]+','-',name.lower()).strip('-')[:45]+'-'+tenant[:6]
    brand={'name':name,'tagline':'Every introduction matters.','description':'','logo':'','logo_scale':100,'background':'#164e40','images':[],'email_domain':'','website':'','allow_employee_banners':True,'analytics_salt':secrets.token_hex(32),'qr':{'foreground':'#164e40','background':'#ffffff','logo':False,'frame':'Scan to connect'}}
    c.execute('INSERT INTO organizations VALUES(?,?,?,?,?,?)',(tenant,name,slug,'active',dump(brand),now()))
    c.execute('INSERT INTO subscriptions(id,tenant_id,plan_id,status,mode,created,updated) VALUES(?,?,?,?,?,?,?)',(uid(),tenant,'free','FREE','none',now(),now()))
    if demo_grant:c.execute('INSERT INTO entitlement_overrides VALUES(?,?,?,?,?,?,?,?,?)',(uid(),tenant,'enterprise','{}','Local demonstration grant','demo',now(),now()+90*86400,None))
    did=uid();c.execute('INSERT INTO divisions VALUES(?,?,?,?,?)',(did,tenant,'General',dump({'description':'','address':'','email':'','phone':'','socials':[]}),now()))
    for i,title in enumerate(['Classic','Portrait','Minimal','Statement','Executive Charcoal','Mint Neon','Vibrant Sunset']):
        tid=uid();policy={'layout':LAYOUTS[i],'accent':({'executive':'#242a29','mint':'#b7f264','sunset':'#ed795e'}.get(LAYOUTS[i],'#164e40')),'visible_fields':FIELDS,'locked_fields':['name','title','email'],'personal_links_allowed':True,'enabled':True,'allowed_blocks':['whatsapp','upi']}
        c.execute('INSERT INTO templates VALUES(?,?,?,?,?)',(tid,tenant,title,dump(policy),now()+i))
        if i==0:brand['default_template']=tid
    c.execute('UPDATE organizations SET data=? WHERE id=?',(dump(brand),tenant))
    add_user(c,tenant,{'name':person_name,'division_id':did},role='owner',account_id=account_id);return tenant

def signup_preferences(c,tenant,d):
    if d.get('handle'):
        handle=text(d['handle'],50).lower()
        if not re.fullmatch(r'[a-z0-9][a-z0-9-]{1,38}[a-z0-9]',handle) or handle in ('api','app','login','signup','dashboard','templates','pricing','features','static','manifest','sw'):raise Problem('Choose a valid card handle.')
        if c.execute('SELECT 1 FROM cards WHERE slug=?',(handle,)).fetchone():raise Problem('This card handle is already claimed.',409)
        c.execute('UPDATE cards SET slug=? WHERE tenant_id=? AND is_default=1',(handle,tenant))
    if d.get('theme') in ('executive','mint','sunset'):
        theme=next((t for t in docs(c,'templates',tenant) if t['data']['layout']==d['theme']),None)
        if theme:c.execute('UPDATE cards SET template_id=? WHERE tenant_id=? AND is_default=1',(theme['id'],tenant))

def initialise(demo=False):
    import migrate
    migrate.apply(sys.modules[__name__])
    with connection() as c:
        if c.execute('SELECT 1 FROM organizations').fetchone():return
        address=os.environ.get('HUDDLE_ADMIN_EMAIL',os.environ.get('HELLO_ADMIN_EMAIL','admin@hello.example' if demo else ''))
        password=os.environ.get('HUDDLE_ADMIN_PASSWORD',os.environ.get('HELLO_ADMIN_PASSWORD','HelloDemo!2026' if demo else ''))
        if not address or not password:raise Problem('Set HELLO_ADMIN_EMAIL and HELLO_ADMIN_PASSWORD or use --demo.')
        aid=uid();c.execute('INSERT INTO accounts VALUES(?,?,?,?,?)',(aid,email(address),password_hash(password),now(),None))
        tenant=create_org(c,os.environ.get('HUDDLE_ORGANIZATION_NAME',os.environ.get('HELLO_COMPANY_NAME','HUDDLE Demo')),aid,'Alex Morgan' if demo else os.environ.get('HUDDLE_ADMIN_NAME',os.environ.get('HELLO_ADMIN_NAME','Workspace owner')),demo_grant=demo)
        if demo:
            did=c.execute('SELECT id FROM divisions WHERE tenant_id=?',(tenant,)).fetchone()[0]
            for name in ['Sophie Chen','Daniel Reed','Maya Patel','James Wilson']:add_user(c,tenant,{'name':name,'email':name.lower().replace(' ','.')+'@hello.example','division_id':did},'HelloDemo!2026')

def card_row(c,slug):
    # Explicit public capability resolution: no collection access by untrusted tenant IDs.
    r=c.execute("SELECT ca.* FROM cards ca JOIN users u ON u.tenant_id=ca.tenant_id AND u.id=ca.user_id JOIN organizations o ON o.id=ca.tenant_id WHERE ca.slug=? AND ca.active=1 AND u.active=1 AND o.status='active'",(slug,)).fetchone()
    if not r:raise Problem('This card is unavailable.',404)
    return r

def public_card(c,slug):
    r=card_row(c,slug);tenant=r['tenant_id'];u=scoped(c,'users',tenant,r['user_id']);t=scoped(c,'templates',tenant,r['template_id']);policy=json.loads(t['data']);brand=company(c,tenant)
    if not policy.get('enabled',True):raise Problem('This card is not published.',404)
    entitlement=billing.effective(sys.modules[__name__],c,tenant)['entitlements']
    if not entitlement['custom_branding']:
        brand={**brand,'logo':'','background':'#164e40'};policy={**policy,'accent':'#164e40'}
    p=json.loads(r['data']);visible=policy['visible_fields'];profile={k:p.get(k,'') for k in visible if k!='banners'}
    if not policy['personal_links_allowed']:profile.pop('personal_links',None)
    division=scoped(c,'divisions',tenant,u['division_id']) if u['division_id'] else None
    banners=p.get('banners',[]) if p.get('banner_mode')=='personal' and brand['allow_employee_banners'] else brand['images']
    if 'banners' not in visible:banners=[]
    blocks=json.loads(r['blocks']);allowed=policy.get('allowed_blocks',['whatsapp','upi']);blocks={k:v for k,v in blocks.items() if k in allowed}
    return {'profile':profile,'company':{k:brand[k] for k in ['name','tagline','description','logo','logo_scale','background','images','website','allow_employee_banners']},'division':dict(name=division['name'],**json.loads(division['data'])) if division else None,'template':{'name':t['name'],'layout':policy['layout'],'accent':policy['accent']},'slug':slug,'card_id':r['id'],'persona':r['persona'],'blocks':blocks,'banner_images':banners,'lead_capture':entitlement['lead_capture'],'show_huddle_branding':not entitlement['remove_branding']}

def edit_card(c,actor,row,data):
    tenant=actor['tenant_id'];admin=actor['role'] in ('owner','admin')
    billing.editable(sys.modules[__name__],c,tenant,row,data)
    if row['user_id']!=actor['id'] and not admin:raise Problem('You can edit only your own cards.',403)
    p=json.loads(row['data']);policy=json.loads(scoped(c,'templates',tenant,row['template_id'])['data']);brand=company(c,tenant)
    for f in FIELDS:
        if f not in data:continue
        v=data[f]
        if not admin and f in policy['locked_fields'] and v!=p.get(f):raise Problem(f'{f} is locked by your organization.',403)
        if f in ('banners','photo'):
            if f=='banners':
                if not isinstance(v,list) or len(v)>3:raise Problem('Use up to three banners.')
                if not admin and not brand['allow_employee_banners'] and v!=p.get(f,[]):raise Problem('Personal banners are disabled.',403)
                v=[asset(c,tenant,x) for x in v]
            else:v=asset(c,tenant,v)
        elif f=='personal_links':
            v=social_links(v)
            if not admin and not policy['personal_links_allowed'] and v!=p.get(f,[]):raise Problem('Personal links are disabled.',403)
        elif f=='email':v=email(v)
        else:v=text(v,1500 if f=='bio' else 300)
        p[f]=v
    if not p.get('name'):raise Problem('Name is required.')
    if 'banner_mode' in data:
        mode=data['banner_mode']
        if mode not in ('company','personal'):raise Problem('Invalid banner mode.')
        if mode!=p.get('banner_mode','company') and not admin and ('banners' in policy['locked_fields'] or not brand['allow_employee_banners']):raise Problem('Banner policy prevents this change.',403)
        p['banner_mode']=mode
    tid=data.get('template_id',row['template_id']);new_policy=json.loads(scoped(c,'templates',tenant,tid)['data'])
    if not new_policy.get('enabled',True):raise Problem('Choose an enabled template.')
    blocks=json.loads(row['blocks'])
    if 'blocks' in data:
        if not isinstance(data['blocks'],dict):raise Problem('Invalid card blocks.')
        blocks={};allowed=new_policy.get('allowed_blocks',['whatsapp','upi'])
        for kind,v in data['blocks'].items():
            if kind not in allowed:raise Problem('This block is disabled by your organization.',403)
            if kind=='whatsapp':
                number=re.sub(r'[ +()-]','',text(v.get('number'),30))
                if not re.fullmatch(r'[1-9][0-9]{7,14}',number):raise Problem('Use an international WhatsApp number with country code.')
                message=text(v.get('message'),500);blocks[kind]={'number':number,'message':message,'url':'https://wa.me/'+number+('?text='+urllib.parse.quote(message) if message else '')}
            elif kind=='upi':
                address=text(v.get('id'),120)
                if not re.fullmatch(r'[a-zA-Z0-9._-]+@[a-zA-Z0-9.-]+',address):raise Problem('Enter a valid UPI ID.')
                blocks[kind]={'id':address,'name':text(v.get('name') or p['name'],100)}
    if data.get('is_default'):
        c.execute('UPDATE cards SET is_default=0 WHERE tenant_id=? AND user_id=?',(tenant,row['user_id']))
    c.execute('UPDATE cards SET data=?,blocks=?,template_id=?,persona=?,is_default=?,active=?,updated=? WHERE tenant_id=? AND id=?',(dump(p),dump(blocks),tid,text(data.get('persona',row['persona']),60),1 if data.get('is_default') else row['is_default'],int(bool(data.get('active',row['active']))),now(),tenant,row['id']))
    publish_assets(c,tenant,[p.get('photo')]+p.get('banners',[]));audit(c,actor,'card.updated',row['id'])

def send_invite(c,tenant,inv):send_invites(c,tenant,[inv])

def send_invites(c,tenant,invitations):
    key=os.environ.get('RESEND_API_KEY');sender=os.environ.get('RESEND_FROM')
    if not invitations or DEMO or not key or not sender or not os.environ.get('PUBLIC_BASE_URL'):return
    brand=company(c,tenant);payload=[]
    for inv in invitations:
        link=security.app_origin(sys.modules[__name__])+'/accept?token='+urllib.parse.quote(inv['token'])
        payload.append({'from':sender,'to':[inv['email']],'subject':'Join '+brand['name']+' on HUDDLE','html':'<p>You are invited to '+html.escape(brand['name'])+'.</p><p><a href="'+html.escape(link,quote=True)+'">Accept invitation</a></p>'})
    batch=hashlib.sha256(':'.join(inv['id'] for inv in invitations).encode()).hexdigest()
    try:
        req=urllib.request.Request('https://api.resend.com/emails/batch',data=dump(payload).encode(),headers={'Authorization':'Bearer '+key,'Content-Type':'application/json','Idempotency-Key':'huddle-'+batch+'-'+str(now()//60)})
        with urllib.request.urlopen(req,timeout=15) as response:result=json.load(response)
        sent=result.get('data',[])
        if len(sent)!=len(invitations) or any(not x.get('id') for x in sent):raise ValueError('Invalid batch result')
        for inv,item in zip(invitations,sent):c.execute("UPDATE invitations SET delivery='sent',provider_id=? WHERE tenant_id=? AND id=?",(item['id'],tenant,inv['id']))
    except (OSError,ValueError):
        for inv in invitations:c.execute("UPDATE invitations SET delivery='failed' WHERE tenant_id=? AND id=?",(tenant,inv['id']))

def invitation(c,actor,data,deliver=True):
    require_admin(actor);tenant=actor['tenant_id'];address=email(data.get('email'));role=data.get('role','member')
    if role not in ('admin','member'):raise Problem('Invite an Admin or Member. Ownership is assigned after joining.')
    if actor['role']!='owner' and role!='member':raise Problem('Only Owners can invite administrators.',403)
    if data.get('division_id'):scoped(c,'divisions',tenant,data['division_id'])
    tid=data.get('template_id') or company(c,tenant)['default_template'];scoped(c,'templates',tenant,tid)
    if not text(data.get('name')):raise Problem('Enter a name.')
    if c.execute('SELECT 1 FROM users u JOIN accounts a ON a.id=u.account_id WHERE u.tenant_id=? AND a.email=?',(tenant,address)).fetchone() or c.execute("SELECT 1 FROM invitations WHERE tenant_id=? AND email=? AND status='pending' AND expires>?",(tenant,address,now())).fetchone():raise Problem('An account or pending invitation already exists.',409)
    billing.capacity(sys.modules[__name__],c,tenant,'seats')
    billing.capacity(sys.modules[__name__],c,tenant,'cards')
    if role=='admin':billing.capacity(sys.modules[__name__],c,tenant,'admins')
    iid=uid();token=secrets.token_urlsafe(32);details={'name':text(data['name']),'title':text(data.get('title')),'division_id':data.get('division_id'),'template_id':tid,'role':role}
    c.execute('INSERT INTO invitations VALUES(?,?,?,?,?,?,?,?,?,?,?)',(iid,tenant,address,dump(details),hashlib.sha256(token.encode()).hexdigest(),token,'pending','disabled' if DEMO else 'pending',now(),now()+7*86400,None))
    inv=dict(c.execute('SELECT * FROM invitations WHERE tenant_id=? AND id=?',(tenant,iid)).fetchone());
    if deliver:send_invite(c,tenant,inv)
    audit(c,actor,'invitation.created',iid);return iid

class Handler(BaseHTTPRequestHandler):
    server_version='HUDDLE/0.7.0'
    def log_message(self,fmt,*args):
        if args and isinstance(args[0],str):args=(args[0].split('?')[0],)+args[1:]
        super().log_message(fmt,*args)
    def send(self,status,content,mime='application/json',headers=None):
        if getattr(self,'_defer_response',False):
            self._pending_response=(status,content,mime,headers);return
        raw=dump(content).encode() if mime=='application/json' or mime.endswith('+json') else content.encode() if isinstance(content,str) else content
        self.send_response(status);self.send_header('Content-Type',mime);self.send_header('Content-Length',str(len(raw)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.send_header('Referrer-Policy','no-referrer')
        self.send_header('Content-Security-Policy',"default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'")
        for k,v in (headers or {}).items():self.send_header(k,v)
        self.end_headers();self.wfile.write(raw)
    def body(self):
        try:n=int(self.headers.get('Content-Length',0))
        except ValueError:raise Problem('Invalid body length.')
        if not 0<=n<=4_400_000:raise Problem('Request too large.',413)
        try:d=json.loads(self.rfile.read(n) or b'{}')
        except (ValueError,UnicodeError):raise Problem('Invalid JSON.')
        if not isinstance(d,dict):raise Problem('Use a JSON object.')
        return d
    def limited(self,group,maximum=12):
        if database.postgres():return database.limited(sys.modules[__name__],self.client_address[0],group,maximum)
        key=(self.client_address[0],group)
        with RATE_LOCK:
            values=[x for x in RATE.get(key,[]) if x>now()-900]
            if len(values)>=maximum:raise Problem('Too many attempts. Try again later.',429)
            RATE[key]=values+[now()]
    def token_hash(self):
        try:cookie=SimpleCookie(self.headers.get('Cookie',''));token=cookie['hello_session'].value
        except (KeyError,ValueError):raise Problem('Please sign in.',401)
        return hashlib.sha256(token.encode()).hexdigest()
    def actor(self,c):
        row=c.execute("SELECT u.*,s.account_id FROM sessions s JOIN users u ON u.account_id=s.account_id AND u.tenant_id=s.active_tenant_id JOIN accounts a ON a.id=s.account_id JOIN organizations o ON o.id=u.tenant_id WHERE s.token_hash=? AND s.expires>? AND u.active=1 AND a.disabled IS NULL AND o.status='active'",(self.token_hash(),now())).fetchone()
        if not row:raise Problem('Please sign in again.',401)
        security.session_valid(sys.modules[__name__],c,self,dict(row))
        expected=self.headers.get('X-Huddle-Tenant')
        if expected and expected!=row['tenant_id']:raise Problem('Organization changed. Refresh before continuing.',409)
        result=user_dict(c,row);result['account_id']=row['account_id'];return result
    def session(self,c,account,tenant,verified=False):
        token=secrets.token_urlsafe(32);c.execute('INSERT INTO sessions(token_hash,account_id,active_tenant_id,expires,mfa_tenant,mfa_at,authenticated_at) VALUES(?,?,?,?,?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),account,tenant,now()+43200,tenant if verified else None,now() if verified else 0,now()))
        return 'hello_session='+token+'; HttpOnly; SameSite=Lax; Path=/; Max-Age=43200'+('; Secure' if base_url().startswith('https://') else '')
    def do_GET(self):self.handle_safe('GET')
    def do_POST(self):self.handle_safe('POST')
    def do_PUT(self):self.handle_safe('PUT')
    def do_DELETE(self):self.handle_safe('DELETE')
    def handle_safe(self,method):
        self._defer_response=True;self._pending_response=None
        try:
            path=urllib.parse.urlsplit(self.path).path
            if method!='GET':
                origin=self.headers.get('Origin');allowed={base_url(),'http://'+self.headers.get('Host',''),'https://'+self.headers.get('Host','')}
                if origin and origin not in allowed:raise Problem('Cross-origin request rejected.',403)
                if 'application/json' not in self.headers.get('Content-Type',''):raise Problem('Use JSON requests.',415)
            with connection() as c:
                c.execute('BEGIN IMMEDIATE' if method!='GET' else 'BEGIN')
                result=self.route(c,method,path)
            if result is not None:self.send(200,result)
        except billing.BillingError as e:self.send(e.code,e.payload)
        except Problem as e:self.send(e.code,{'error':e.message})
        except sqlite3.IntegrityError:self.send(409,{'error':'Duplicate record or invalid organization relationship.'})
        except Exception as e:
            if database.retryable(e):self.send(409,{'error':'Another update completed at the same time. Please retry.'});return
            print('Request failed:',type(e).__name__,flush=True);self.send(500,{'error':'The request could not be completed.'})
        finally:
            self._defer_response=False
            if self._pending_response is not None:self.send(*self._pending_response)
    def query(self):return urllib.parse.parse_qs(urllib.parse.urlsplit(self.path).query)
    def route(self,c,method,path):
        import sales
        matched,result=sales.routes(sys.modules[__name__],self,c,method,path)
        if matched:return result
        matched,result=platform_admin.route(sys.modules[__name__],self,c,method,path)
        if matched:return result
        matched,result=billing.routes(sys.modules[__name__],self,c,method,path)
        if matched:return result
        matched,result=security.handle(sys.modules[__name__],self,c,method,path)
        if matched:return result
        matched,result=surfaces.serve(sys.modules[__name__],self,c,method,path)
        if matched:return result
        if path=='/api/health' and method=='GET':return {'ok':True,'application':'HUDDLE','schema_version':3,'version':'0.7.0'}
        if path=='/api/setup' and method=='GET':return {'company_name':'HUDDLE','demo':DEMO,'email_configured':not DEMO and bool(os.environ.get('RESEND_API_KEY') and os.environ.get('RESEND_FROM') and os.environ.get('PUBLIC_BASE_URL')),'join_enabled':False,'divisions':[],'multi_tenant':True,'registration_enabled':True,'google_configured':security.google_configured()}
        if path=='/api/register' and method=='POST':
            self.limited('register',8);d=self.body();aid=uid();c.execute('INSERT INTO accounts VALUES(?,?,?,?,?)',(aid,email(d.get('email')),password_hash(d.get('password')),now(),None));tenant=create_org(c,d.get('organization'),aid,text(d.get('name')))
            signup_preferences(c,tenant,d)
            self.send(200,{'ok':True},headers={'Set-Cookie':self.session(c,aid,tenant)});return
        if path=='/api/login' and method=='POST':
            self.limited('login');d=self.body();a=c.execute('SELECT * FROM accounts WHERE email=? AND disabled IS NULL',(email(d.get('email')),)).fetchone()
            if not a or not password_matches(str(d.get('password','')),a['password_hash']):raise Problem('Email or password is incorrect.',401)
            membership=c.execute("SELECT u.* FROM users u JOIN organizations o ON o.id=u.tenant_id WHERE u.account_id=? AND u.active=1 AND o.status='active' ORDER BY u.created",(a['id'],)).fetchone()
            if not membership:raise Problem('No active organization membership.',403)
            if d.get('tenant_id'):
                membership=c.execute("SELECT u.* FROM users u JOIN organizations o ON o.id=u.tenant_id WHERE u.account_id=? AND u.tenant_id=? AND u.active=1 AND o.status='active'",(a['id'],d['tenant_id'])).fetchone()
                if not membership:raise Problem('No active membership.',403)
            if security.enabled(c,membership['tenant_id'],membership['id']):return security.challenge(sys.modules[__name__],c,membership)
            self.send(200,{'user':user_dict(c,membership)},headers={'Set-Cookie':self.session(c,a['id'],membership['tenant_id'])});return
        if path=='/api/logout' and method=='POST':
            try:c.execute('DELETE FROM sessions WHERE token_hash=?',(self.token_hash(),))
            except Problem:pass
            self.send(200,{'ok':True},headers={'Set-Cookie':'hello_session=; HttpOnly; SameSite=Lax; Path=/; Max-Age=0'});return
        if path in ('/api/invitation','/api/accept'):
            d=self.body() if method=='POST' else {};token=d.get('token') or self.query().get('token',[''])[0]
            # Unforgeable invitation token resolves its tenant before any membership operation.
            inv=c.execute('SELECT * FROM invitations WHERE token_hash=?',(hashlib.sha256(str(token).encode()).hexdigest(),)).fetchone()
            if not inv or inv['status']!='pending' or inv['expires']<now():raise Problem('Invitation expired or unavailable.',410)
            tenant=inv['tenant_id'];details=json.loads(inv['data'])
            if path=='/api/invitation' and method=='GET':
                try:signed_in=self.actor(c)['account_id']==c.execute('SELECT id FROM accounts WHERE email=?',(inv['email'],)).fetchone()[0]
                except (Problem,TypeError):signed_in=False
                return {'name':details['name'],'email':inv['email'],'company_name':company(c,tenant)['name'],'signed_in':signed_in}
            if path!='/api/accept' or method!='POST':raise Problem('Route not found.',404)
            self.limited('accept',20);existing=c.execute('SELECT * FROM accounts WHERE email=?',(inv['email'],)).fetchone()
            if existing:
                try:signed_in=self.actor(c)['account_id']==existing['id']
                except Problem:signed_in=False
                if not signed_in and not password_matches(str(d.get('password','')),existing['password_hash']):raise Problem('Sign in to your existing account or confirm its password to join this organization.',403)
                add_user(c,tenant,{**details,'name':text(d.get('name')) or details['name']},role=details['role'],account_id=existing['id'])
            else:add_user(c,tenant,{**details,'email':inv['email'],'name':text(d.get('name')) or details['name']},d.get('password'),details['role'])
            c.execute("UPDATE invitations SET status='accepted' WHERE tenant_id=? AND id=? AND status='pending'",(tenant,inv['id']));return {'ok':True}
        if path=='/api/join' and method=='POST':
            self.limited('join',8);d=self.body();org=c.execute("SELECT id FROM organizations WHERE slug=? AND status='active'",(text(d.get('organization_slug')),)).fetchone()
            if not org:raise Problem('Organization not found.',404)
            tenant=org['id'];brand=company(c,tenant);address=email(d.get('email'))
            if not brand.get('email_domain') or address.split('@')[1]!=brand['email_domain'].lower():raise Problem('Ask your administrator for an invitation.',403)
            if not text(d.get('name')):raise Problem('Enter your name.')
            c.execute('INSERT INTO requests VALUES(?,?,?,?,?,?,?)',(uid(),tenant,address,dump({'name':text(d['name']),'title':text(d.get('title'))}),password_hash(d.get('password')),'pending',now()));return {'ok':True}
        if method=='GET' and path.startswith('/api/card/'):return public_card(c,path.rsplit('/',1)[-1])
        if method=='GET' and path.startswith('/n/'):
            r=c.execute("SELECT n.*,ca.slug FROM nfc_cards n JOIN cards ca ON ca.tenant_id=n.tenant_id AND ca.id=n.card_id WHERE n.token=? AND n.active=1",(path.rsplit('/',1)[-1],)).fetchone()
            if not r:raise Problem('NFC card unavailable.',404)
            public_card(c,r['slug']);self.send(302,b'','text/plain',{'Location':'/c/'+r['slug']+'?via=nfc'});return
        if method=='POST' and path=='/api/activity':
            self.limited('activity',120);d=self.body();r=card_row(c,text(d.get('slug')));public_card(c,r['slug'])
            if d.get('kind') not in ('view','link_click'):raise Problem('Invalid event.')
            cookie=self.record_activity(c,r,d['kind'],d.get('source','link'),text(d.get('label'),40));self.send(200,{'ok':True},headers={'Set-Cookie':cookie} if cookie else None);return
        if method=='POST' and path.startswith('/api/exchange/'):
            self.limited('leads',30);r=card_row(c,path.rsplit('/',1)[-1]);public_card(c,r['slug']);billing.require(sys.modules[__name__],c,r['tenant_id'],'lead_capture');d=self.body()
            if d.get('website_trap'):return {'ok':True}
            if d.get('consent') is not True:raise Problem('Consent is required to share your details.')
            data={'name':text(d.get('name'),150),'email':email(d.get('email')),'phone':text(d.get('phone'),40),'company':text(d.get('company'),150),'message':text(d.get('message'),2000)}
            if not data['name']:raise Problem('Enter your name.')
            cid=uid();c.execute('INSERT INTO contact_exchanges VALUES(?,?,?,?,?,?,?,?,?)',(cid,r['tenant_id'],r['user_id'],r['id'],dump(data),'huddle-exchange-v1',now(),text(d.get('source') or 'card',20),now()));return {'ok':True}
        if method=='GET' and (path.startswith('/contact/') or path.startswith('/qr/') or path.startswith('/upi/')):
            slug=path.rsplit('/',1)[-1].rsplit('.',1)[0];r=card_row(c,slug);card=public_card(c,slug)
            if path.startswith('/contact/'):
                p=card['profile'];esc=lambda v:str(v or '').replace('\\','\\\\').replace('\n','\\n').replace('\r','').replace(';','\\;').replace(',','\\,')
                lines=['BEGIN:VCARD','VERSION:3.0','FN:'+esc(p.get('name')),'ORG:'+esc(card['company']['name']),'TITLE:'+esc(p.get('title'))]
                if p.get('email'):lines.append('EMAIL;TYPE=WORK:'+esc(p['email']))
                if p.get('phone'):lines.append('TEL;TYPE=WORK:'+esc(p['phone']))
                lines+=['URL:'+base_url()+'/c/'+slug,'END:VCARD'];folded=[]
                for line in lines:
                    buf=''
                    for char in line:
                        if len((buf+char).encode())>72:folded.append(buf);buf=' '
                        buf+=char
                    folded.append(buf)
                cookie=self.record_activity(c,r,'contact_download',self.query().get('via',['link'])[0]);self.send(200,'\r\n'.join(folded)+'\r\n','text/vcard; charset=utf-8',{'Content-Disposition':'attachment; filename="'+slug+'.vcf"',**({'Set-Cookie':cookie} if cookie else {})});return
            target=base_url()+'/c/'+slug+'?via=qr';qrconfig=company(c,r['tenant_id']).get('qr',{}) if billing.effective(sys.modules[__name__],c,r['tenant_id'])['entitlements']['custom_qr'] else {}
            if path.startswith('/upi/'):
                block=card['blocks'].get('upi')
                if not block:raise Problem('UPI block unavailable.',404)
                target='upi://pay?'+urllib.parse.urlencode({'pa':block['id'],'pn':block['name'],'cu':'INR'});qrconfig={}
            svg=self.qr_svg(c,r['tenant_id'],target,qrconfig)
            if path.endswith('.png'):self.send(200,self.qr_png(svg),'image/png')
            else:self.send(200,svg,'image/svg+xml')
            return
        if path.startswith('/api/'):
            actor=self.actor(c);tenant=actor['tenant_id'];admin=actor['role'] in ('owner','admin')
            if (method!='GET' or path.endswith('/export')) and path not in ('/api/share',) and not security.recent(sys.modules[__name__],c,self,actor):
                self.send(403,{'error':'Verify your authenticator to continue.','elevation_required':True});return
            if path=='/api/organizations' and method=='GET':return {'organizations':[dict(r) for r in c.execute("SELECT o.id,o.name,o.slug,u.role FROM users u JOIN organizations o ON o.id=u.tenant_id WHERE u.account_id=? AND u.active=1 AND o.status='active'",(actor['account_id'],))],'active_tenant_id':tenant,'ownership':billing.organization_capacity(sys.modules[__name__],c,actor['account_id'])}
            if path=='/api/organizations' and method=='POST':
                d=self.body();new=create_org(c,d.get('name'),actor['account_id'],actor.get('name'));return {'id':new}
            if path=='/api/organizations/switch' and method=='POST':
                d=self.body();target=d.get('tenant_id');m=c.execute("SELECT 1 FROM users u JOIN organizations o ON o.id=u.tenant_id WHERE u.tenant_id=? AND u.account_id=? AND u.active=1 AND o.status='active'",(target,actor['account_id'])).fetchone()
                if not m:raise Problem('No active membership in this organization.',403)
                member=c.execute('SELECT * FROM users WHERE tenant_id=? AND account_id=? AND active=1',(target,actor['account_id'])).fetchone()
                if security.enabled(c,target,member['id']):return security.challenge(sys.modules[__name__],c,member,'switch',{'session':self.token_hash()})
                c.execute('UPDATE sessions SET active_tenant_id=? WHERE token_hash=? AND account_id=?',(target,self.token_hash(),actor['account_id']));return {'ok':True}
            if path=='/api/state' and method=='GET':
                brand=company(c,tenant);brand={k:v for k,v in brand.items() if k!='analytics_salt'};result={'user':{k:v for k,v in actor.items() if k!='account_id'},'company':brand,'organization':{'id':tenant,'name':brand['name']},'divisions':docs(c,'divisions',tenant),'templates':docs(c,'templates',tenant),'invitation_capacity':billing.invitation_capacity(sys.modules[__name__],c,tenant) if admin else None,'base_url':base_url(),'platform_access':bool(c.execute('SELECT 1 FROM platform_admin_grants WHERE account_id=?',(actor['account_id'],)).fetchone()),'cards':self.cards(c,actor),'billing':billing.summary(sys.modules[__name__],c,tenant) if actor['role']=='owner' else {k:v for k,v in billing.effective(sys.modules[__name__],c,tenant).items() if k!='subscription'}}
                if admin:
                    result['team']=[user_dict(c,r) for r in c.execute('SELECT * FROM users WHERE tenant_id=? ORDER BY created',(tenant,))]
                    result['invitations']=[{k:r[k] for k in ('id','email','status','delivery','created','expires')}|{'data':json.loads(r['data']),'url':security.app_origin(sys.modules[__name__])+'/accept?token='+r['token']} for r in c.execute('SELECT * FROM invitations WHERE tenant_id=? ORDER BY created DESC',(tenant,)) if actor['role']=='owner' or json.loads(r['data']).get('role')=='member']
                    result['requests']=[{k:r[k] for k in ('id','email','status','created')}|{'data':json.loads(r['data'])} for r in c.execute('SELECT * FROM requests WHERE tenant_id=? ORDER BY created DESC',(tenant,))]
                    result['audit']=[dict(r) for r in c.execute('SELECT * FROM audit WHERE tenant_id=? ORDER BY id DESC LIMIT 40',(tenant,))]
                return result
            if path=='/api/password' and method=='POST':
                if not security.enabled(c,tenant,actor['id']) and c.execute('SELECT 1 FROM user_security s JOIN users u ON u.tenant_id=s.tenant_id AND u.id=s.user_id WHERE u.account_id=? AND s.enabled=1 AND u.active=1',(actor['account_id'],)).fetchone():raise Problem('Switch to your authenticator-protected organization before changing the account password.',403)
                d=self.body();a=c.execute('SELECT * FROM accounts WHERE id=?',(actor['account_id'],)).fetchone()
                if not password_matches(str(d.get('current','')),a['password_hash']):raise Problem('Current password is incorrect.',403)
                c.execute('UPDATE accounts SET password_hash=? WHERE id=?',(password_hash(d.get('password')),actor['account_id']));c.execute('DELETE FROM sessions WHERE account_id=?',(actor['account_id'],));c.execute('DELETE FROM platform_sessions WHERE account_id=?',(actor['account_id'],));return {'ok':True}
            if path=='/api/cards' and method=='GET':return {'cards':self.cards(c,actor)}
            if path=='/api/cards' and method=='POST':
                d=self.body();member_id=d.get('user_id',actor['id'])
                if member_id!=actor['id']:require_admin(actor)
                m=scoped(c,'users',tenant,member_id)
                if c.execute('SELECT count(*) FROM cards WHERE tenant_id=? AND user_id=?',(tenant,member_id)).fetchone()[0]>=20:raise Problem('Use up to twenty personas.')
                source=c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? AND is_default=1',(tenant,member_id)).fetchone();profile={**json.loads(source['data']),'template_id':source['template_id']}
                cid=create_card(c,tenant,member_id,profile,d.get('persona','Community'));edit_card(c,actor,scoped(c,'cards',tenant,cid),d);return {'id':cid}
            if path.startswith('/api/cards/') and method=='PUT':
                r=scoped(c,'cards',tenant,path.rsplit('/',1)[-1]);edit_card(c,actor,r,self.body());return {'ok':True}
            if path.startswith('/api/cards/') and path.endswith('/wallet/apple') or path.startswith('/api/cards/') and path.endswith('/wallet/google'):
                if method!='POST':raise Problem('Use POST.',405)
                cid=path.split('/')[3];r=scoped(c,'cards',tenant,cid)
                if r['user_id']!=actor['id']:require_admin(actor)
                billing.require(sys.modules[__name__],c,tenant,'wallet')
                provider=path.rsplit('/',1)[-1];self.send(503,{'error':'Wallet provider configuration required.','provider':provider,'status':'configuration_required','required':['Apple pass certificate, team ID and pass type ID'] if provider=='apple' else ['Google Wallet issuer ID and service account']});return
            if path.startswith('/api/cards/') and (path.endswith('/signature') or path.endswith('/background.png')) and method=='GET':
                row=scoped(c,'cards',tenant,path.split('/')[3])
                if row['user_id']!=actor['id']:require_admin(actor)
                model=public_card(c,row['slug']);profile=model['profile'];brand=model['company'];url=base_url()+'/c/'+row['slug']
                if path.endswith('/signature'):
                    body='<table style="font-family:Arial,sans-serif;color:#164e40"><tr><td style="padding:16px"><strong>'+html.escape(profile.get('name',''))+'</strong><br>'+html.escape(profile.get('title',''))+'<br>'+html.escape(brand['name'])+'<br><a href="'+html.escape(url,quote=True)+'">Connect with me</a></td><td><img width="110" alt="Scan to connect" src="'+html.escape(base_url()+'/qr/'+row['slug']+'.png',quote=True)+'"></td></tr></table>'
                    self.send(200,body,'text/html; charset=utf-8',{'Content-Disposition':'attachment; filename="huddle-signature.html"'});return
                from PIL import Image,ImageDraw,ImageFont
                image=Image.new('RGB',(1920,1080),brand['background']);draw=ImageDraw.Draw(image);draw.ellipse((800,-500,2100,800),fill='#233e34');draw.text((100,800),profile.get('name',''),fill='white',font=ImageFont.load_default(size=64));draw.text((100,900),brand['name'],fill='white',font=ImageFont.load_default(size=40));qr=Image.open(io.BytesIO(self.qr_png(self.qr_svg(c,tenant,url+'?via=qr',{})))).resize((310,310));image.paste(qr,(1500,670));stream=io.BytesIO();image.save(stream,format='PNG');self.send(200,stream.getvalue(),'image/png',{'Content-Disposition':'attachment; filename="huddle-meeting-background.png"'});return
            if path=='/api/upload' and method=='POST':return self.upload(c,actor,self.body())
            if path.startswith('/api/users/') and method in ('PUT','DELETE'):
                r=scoped(c,'users',tenant,path.rsplit('/',1)[-1]);own=r['id']==actor['id'];d=self.body()
                if not own:require_admin(actor)
                if 'role' in d and d['role']!=r['role'] or method=='DELETE' and r['role']!='member':
                    require_owner(actor)
                    if own:raise Problem('You cannot remove or demote yourself.')
                    if r['role']=='owner' and c.execute("SELECT count(*) FROM users WHERE tenant_id=? AND role='owner' AND active=1",(tenant,)).fetchone()[0]<=1:raise Problem('The last Owner must remain active.')
                if method=='DELETE':
                    if own:raise Problem('You cannot remove yourself.')
                    c.execute('UPDATE users SET active=0,disabled=? WHERE tenant_id=? AND id=?',(now(),tenant,r['id']));audit(c,actor,'member.removed',r['id']);return {'ok':True}
                if not own and r['role']=='owner' and actor['role']!='owner':raise Problem('Only Owners can edit another Owner.',403)
                if 'role' in d:
                    if d['role'] not in ('owner','admin','member'):raise Problem('Invalid role.')
                    if d['role']!=r['role']:
                        if r['role']=='member' and d['role'] in ('owner','admin'):billing.capacity(sys.modules[__name__],c,tenant,'admins')
                        c.execute('UPDATE users SET role=? WHERE tenant_id=? AND id=?',(d['role'],tenant,r['id']))
                if 'division_id' in d:
                    if d['division_id']!=r['division_id']:require_admin(actor)
                    if d['division_id']:scoped(c,'divisions',tenant,d['division_id'])
                    c.execute('UPDATE users SET division_id=? WHERE tenant_id=? AND id=?',(d['division_id'] or None,tenant,r['id']))
                card=c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? AND is_default=1',(tenant,r['id'])).fetchone();edit_card(c,actor,card,d);return {'ok':True}
            if path=='/api/share' and method=='POST':
                d=self.body()
                if d.get('card_id'):r=scoped(c,'cards',tenant,d['card_id'])
                else:r=c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? AND is_default=1',(tenant,d.get('user_id',actor['id']))).fetchone()
                if not r:raise Problem('Card not found.',404)
                if r['user_id']!=actor['id']:require_admin(actor)
                public_card(c,r['slug']);self.record_activity(c,r,'share','staff');return {'url':base_url()+'/c/'+r['slug']}
            if path in ('/api/leads','/api/leads/export') and method=='GET':return self.leads(c,actor,path.endswith('/export'))
            if path in ('/api/analytics','/api/analytics/export') and method=='GET':return self.analytics(c,actor,path.endswith('/export'))
            require_admin(actor)
            if path=='/api/company' and method=='PUT':return self.branding(c,actor,self.body())
            if (path=='/api/divisions' or path.startswith('/api/divisions/')) and method in ('POST','PUT'):
                d=self.body();rid=uid() if method=='POST' else path.rsplit('/',1)[-1]
                if method=='PUT':scoped(c,'divisions',tenant,rid)
                value={k:text(d.get(k),1000) for k in ('description','address','phone','email')};value['socials']=social_links(d.get('socials',[]))
                if method=='POST':c.execute('INSERT INTO divisions VALUES(?,?,?,?,?)',(rid,tenant,text(d.get('name')) or 'General',dump(value),now()))
                else:c.execute('UPDATE divisions SET name=?,data=? WHERE tenant_id=? AND id=?',(text(d.get('name')) or 'General',dump(value),tenant,rid))
                return {'id':rid}
            if (path=='/api/templates' or path.startswith('/api/templates/')) and method in ('POST','PUT'):return self.template(c,actor,self.body(),path.rsplit('/',1)[-1] if method=='PUT' else None)
            if path=='/api/invitations' and method=='POST':return {'id':invitation(c,actor,self.body())}
            if path.startswith('/api/invitations/') and method in ('POST','DELETE'):
                r=scoped(c,'invitations',tenant,path.rsplit('/',1)[-1])
                if actor['role']!='owner' and json.loads(r['data']).get('role')!='member':raise Problem('Owner access required.',403)
                if method=='DELETE':c.execute("UPDATE invitations SET status='revoked' WHERE tenant_id=? AND id=?",(tenant,r['id']))
                else:
                    if r['status']!='pending' or r['expires']<now():raise Problem('Invitation unavailable.')
                    send_invite(c,tenant,dict(r))
                return {'ok':True}
            if path.startswith('/api/requests/') and method=='POST':
                r=scoped(c,'requests',tenant,path.rsplit('/',1)[-1]);d=self.body()
                if r['status']!='pending':raise Problem('Request already reviewed.')
                if d.get('approve'):
                    details=json.loads(r['data']);existing=c.execute('SELECT * FROM accounts WHERE email=?',(r['email'],)).fetchone()
                    if existing:raise Problem('Existing account: send an invitation so the account holder can authenticate.')
                    did=c.execute('SELECT id FROM divisions WHERE tenant_id=? ORDER BY created LIMIT 1',(tenant,)).fetchone()[0];mid=add_user(c,tenant,{**details,'email':r['email'],'division_id':did},secrets.token_urlsafe(20));account=scoped(c,'users',tenant,mid)['account_id'];c.execute('UPDATE accounts SET password_hash=? WHERE id=?',(r['password'],account))
                c.execute('UPDATE requests SET status=? WHERE tenant_id=? AND id=?',('approved' if d.get('approve') else 'rejected',tenant,r['id']));return {'ok':True}
            if path.startswith('/api/import/'):return self.onboarding(c,actor,method,path)
            if path.startswith('/api/nfc/'):return self.nfc(c,actor,method,path)
            raise Problem('Route not found.',404)
        if method=='GET':
            if path=='/downloads/hello-admin-debug.apk' or path=='/downloads/huddle-admin-debug.apk':self.send(200,(ROOT/'downloads/hello-admin-debug.apk').read_bytes(),'application/vnd.android.package-archive');return
            if path.startswith('/static/'):
                p=(ROOT/path.lstrip('/')).resolve()
                if (ROOT/'static').resolve() not in p.parents or not p.is_file():raise Problem('Not found.',404)
                self.send(200,p.read_bytes(),mimetypes.guess_type(p)[0] or 'application/octet-stream');return
            if path.startswith('/media/'):
                r=c.execute("SELECT tenant_id,visibility FROM assets WHERE storage_key=?",(path,)).fetchone()
                if not r:raise Problem('Image unavailable.',404)
                if r['visibility']!='public':
                    if self.actor(c)['tenant_id']!=r['tenant_id']:raise Problem('Image unavailable.',404)
                try:raw=object_storage.read(sys.modules[__name__],path)
                except (OSError,ValueError):raise Problem('Image unavailable.',404)
                self.send(200,raw,'image/png');return
            raise Problem('Page not found.',404)
        raise Problem('Route not found.',404)
    def cards(self,c,actor):
        # Members never receive another member's persona collection.
        return [{**dict(r),'data':json.loads(r['data']),'blocks':json.loads(r['blocks'])} for r in c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? ORDER BY is_default DESC,created',(actor['tenant_id'],actor['id']))]
    def upload(self,c,actor,d):
        if d.get('purpose') not in ('photo','personal'):require_admin(actor)
        try:raw=base64.b64decode(d.get('data',''),validate=True)
        except ValueError:raise Problem('Invalid image.')
        if len(raw)>3*1024*1024:raise Problem('Use an image smaller than 3 MB.')
        from PIL import Image,ImageOps
        try:
            im=Image.open(io.BytesIO(raw))
            if im.format not in ('PNG','JPEG','WEBP') or im.width*im.height>25000000:raise ValueError()
            im=ImageOps.exif_transpose(im).convert('RGBA');im.thumbnail((1600,1600));stream=io.BytesIO();im.save(stream,format='PNG');raw=stream.getvalue()
        except Exception:raise Problem('Use a valid PNG, JPEG or WebP up to 25 megapixels.')
        if len(raw)>3*1024*1024:raise Problem('The processed image is too large. Resize it and try again.')
        tenant=actor['tenant_id'];name=hashlib.sha256(raw).hexdigest()+'.png';url='/media/'+tenant+'/'+name
        object_storage.write(sys.modules[__name__],url,raw)
        c.execute('INSERT OR IGNORE INTO assets VALUES(?,?,?,?,?,?)',(uid(),tenant,actor['id'],url,'private',now()));return {'url':url}
    def branding(self,c,actor,d):
        tenant=actor['tenant_id'];brand=company(c,tenant)
        if any(k in d and d[k]!=brand.get(k) for k in ('background','logo','logo_scale')):billing.require(sys.modules[__name__],c,tenant,'custom_branding')
        if 'qr' in d and d['qr']!=brand.get('qr'):billing.require(sys.modules[__name__],c,tenant,'custom_qr')
        for k in ('name','tagline','description','website','email_domain'):
            if k in d:brand[k]=text(d[k],1500 if k=='description' else 300)
        if not brand['name']:raise Problem('Organization name is required.')
        if brand['website'] and not re.fullmatch(r'https?://[^\s]+',brand['website']):raise Problem('Use a full website URL.')
        if brand['email_domain'] and not re.fullmatch(r'[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}',brand['email_domain']):raise Problem('Use an email domain without @.')
        if 'background' in d:brand['background']=colour(d['background'])
        if 'logo' in d:brand['logo']=asset(c,tenant,d['logo'])
        try:brand['logo_scale']=max(40,min(180,int(d.get('logo_scale',brand['logo_scale']))))
        except (ValueError,TypeError):raise Problem('Invalid logo scale.')
        if 'images' in d:
            if not isinstance(d['images'],list) or len(d['images'])>3:raise Problem('Use up to three company banners.')
            brand['images']=[asset(c,tenant,x) for x in d['images']]
        if 'default_template' in d:
            t=scoped(c,'templates',tenant,d['default_template'])
            if not json.loads(t['data']).get('enabled',True):raise Problem('Choose an enabled default template.')
            brand['default_template']=t['id']
        if 'allow_employee_banners' in d:brand['allow_employee_banners']=bool(d['allow_employee_banners'])
        if 'qr' in d:
            q=d['qr'];fg=colour(q.get('foreground','#164e40'));bg=colour(q.get('background','#ffffff'))
            def lum(v):
                rgb=[int(v[i:i+2],16)/255 for i in (1,3,5)];rgb=[x/12.92 if x<=.04045 else ((x+.055)/1.055)**2.4 for x in rgb];return sum(x*w for x,w in zip(rgb,[.2126,.7152,.0722]))
            if (lum(bg)+.05)/(lum(fg)+.05)<4.5:raise Problem('QR codes need dark modules on a light background with at least 4.5:1 contrast.')
            eye=q.get('eye_style','square')
            if eye not in ('square','rounded'):raise Problem('Invalid QR eye style.')
            brand['qr']={'foreground':fg,'background':bg,'logo':bool(q.get('logo')),'frame':text(q.get('frame'),60),'eye_style':eye}
        c.execute('UPDATE organizations SET name=?,data=? WHERE id=?',(brand['name'],dump(brand),tenant));publish_assets(c,tenant,[brand['logo']]+brand['images']);audit(c,actor,'organization.updated',tenant);return {'ok':True}
    def template(self,c,actor,d,rid):
        tenant=actor['tenant_id']
        if rid:scoped(c,'templates',tenant,rid)
        if d.get('accent','#164e40')!='#164e40':billing.require(sys.modules[__name__],c,tenant,'custom_branding')
        if d.get('layout') not in LAYOUTS or not text(d.get('name')):raise Problem('Choose a name and valid layout.')
        visible=d.get('visible_fields',FIELDS);locked=d.get('locked_fields',[])
        if not isinstance(visible,list) or not isinstance(locked,list) or any(k not in FIELDS for k in visible+locked):raise Problem('Invalid fields.')
        blocks=d.get('allowed_blocks',['whatsapp','upi'])
        if not isinstance(blocks,list) or any(k not in ('whatsapp','upi') for k in blocks):raise Problem('Invalid allowed blocks.')
        policy={'layout':d['layout'],'accent':colour(d.get('accent','#164e40')),'visible_fields':visible,'locked_fields':locked,'personal_links_allowed':bool(d.get('personal_links_allowed',True)),'enabled':bool(d.get('enabled',True)),'allowed_blocks':blocks}
        # Organization-wide restrictions prevent bypass by selecting another persona/template.
        for r in c.execute('SELECT id,data FROM templates WHERE tenant_id=?',(tenant,)).fetchall():
            old=json.loads(r['data']);old.update(locked_fields=locked,personal_links_allowed=policy['personal_links_allowed'],allowed_blocks=blocks);c.execute('UPDATE templates SET data=? WHERE tenant_id=? AND id=?',(dump(old),tenant,r['id']))
        if rid:c.execute('UPDATE templates SET name=?,data=? WHERE tenant_id=? AND id=?',(text(d['name']),dump(policy),tenant,rid))
        else:rid=uid();c.execute('INSERT INTO templates VALUES(?,?,?,?,?)',(rid,tenant,text(d['name']),dump(policy),now()))
        audit(c,actor,'template.updated',rid);return {'id':rid}
    def onboarding(self,c,actor,method,path):
        tenant=actor['tenant_id'];billing.require(sys.modules[__name__],c,tenant,'bulk_import');divisions=docs(c,'divisions',tenant)
        if method=='GET' and path=='/api/import/template':
            if self.query().get('format',[''])[0]=='csv':
                stream=io.StringIO();writer=csv.writer(stream);writer.writerow(['Name','Email','Job title','Division']);writer.writerow(['Example Person','person@example.com','Sales',divisions[0]['name']]);self.send(200,stream.getvalue(),'text/csv; charset=utf-8',{'Content-Disposition':'attachment; filename="huddle-team.csv"'});return
            import openpyxl
            b=openpyxl.Workbook();s=b.active;s.title='Team invitations';s.append(['Name','Email','Job title','Division']);s.append(['Example Person','person@example.com','Sales',divisions[0]['name']]);out=io.BytesIO();b.save(out);self.send(200,out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',{'Content-Disposition':'attachment; filename="huddle-team.xlsx"'});return
        if method=='POST' and path=='/api/import/preview':
            d=self.body()
            try:
                raw=base64.b64decode(d.get('data',''),validate=True)
                if len(raw)>2*1024*1024:raise ValueError()
                if d.get('format')=='csv' or str(d.get('filename','')).lower().endswith('.csv'):
                    reader=csv.reader(io.StringIO(raw.decode('utf-8-sig')));rows=[]
                    for i,row in enumerate(reader):
                        if i>101:raise ValueError()
                        rows.append(row)
                else:
                    import openpyxl
                    with zipfile.ZipFile(io.BytesIO(raw)) as z:
                        if sum(i.file_size for i in z.infolist())>20*1024*1024:raise ValueError()
                    b=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True);rows=list(b.active.iter_rows(max_row=103,max_col=4,values_only=True));b.close()
                if not rows or list(rows[0])!=['Name','Email','Job title','Division']:raise ValueError()
                entries=[r for r in rows[1:] if any(v is not None and v!='' for v in r)]
                if len(entries)>100:raise ValueError()
            except Exception:raise Problem('Use the supplied CSV/XLSX template, up to 100 members and 2 MB.')
            known={v['name'].lower():v['id'] for v in divisions};seen=set();result=[]
            for i,row in enumerate(entries,2):
                error='';r=(list(row)+['']*4)[:4];name,address,title,division=[text(v) for v in r]
                try:
                    address=email(address)
                    if not name:raise Problem('Name required.')
                    if division.lower() not in known:raise Problem('Division must belong to this organization.')
                    if address in seen or c.execute('SELECT 1 FROM users u JOIN accounts a ON a.id=u.account_id WHERE u.tenant_id=? AND a.email=?',(tenant,address)).fetchone():raise Problem('Duplicate member.')
                except Problem as e:error=e.message
                seen.add(address);result.append({'row':i,'name':name,'email':address,'title':title,'division':division,'division_id':known.get(division.lower(),''),'error':error})
            return {'rows':result}
        if method=='POST' and path=='/api/import/confirm':
            rows=self.body().get('rows')
            if not isinstance(rows,list) or not 1<=len(rows)<=100:raise Problem('Import 1–100 rows.')
            results=[]
            for i,row in enumerate(rows):
                # Savepoints keep a rejected row from leaving partial changes behind.
                c.execute('SAVEPOINT import_row')
                try:results.append({'email':row.get('email'),'id':invitation(c,actor,{**row,'role':'member'},deliver=False)});c.execute('RELEASE import_row')
                except (Problem,sqlite3.IntegrityError) as e:
                    c.execute('ROLLBACK TO import_row');c.execute('RELEASE import_row');results.append({'email':row.get('email'),'error':e.message if isinstance(e,Problem) else 'Duplicate invitation.'})
            pending=[dict(scoped(c,'invitations',tenant,r['id'])) for r in results if 'id' in r]
            send_invites(c,tenant,pending)
            return {'results':results}
        raise Problem('Route not found.',404)
    def nfc(self,c,actor,method,path):
        tenant=actor['tenant_id']
        if method=='GET' and path=='/api/nfc/targets':
            targets=[]
            for u in c.execute('SELECT * FROM users WHERE tenant_id=? AND active=1',(tenant,)):
                d=user_dict(c,u);targets.append({k:d.get(k) for k in ('id','name','division_id','slug','card_id')})
            cards=[{**dict(r),'url':base_url()+'/n/'+r['token']} for r in c.execute('SELECT n.*,ca.user_id FROM nfc_cards n JOIN cards ca ON ca.tenant_id=n.tenant_id AND ca.id=n.card_id WHERE n.tenant_id=? ORDER BY n.created DESC',(tenant,))]
            return {'targets':targets,'cards':cards,'base_url':base_url(),'android_apk':'/downloads/huddle-admin-debug.apk'}
        d=self.body()
        if method=='POST' and path=='/api/nfc/issue':
            if d.get('card_id'):r=scoped(c,'cards',tenant,d['card_id'])
            else:r=c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? AND is_default=1',(tenant,d.get('user_id'))).fetchone()
            if not r:raise Problem('Choose a card in this organization.',404)
            public_card(c,r['slug']);iid=uid();token=secrets.token_urlsafe(12);c.execute('INSERT INTO nfc_cards VALUES(?,?,?,?,?,?,?,?,?)',(iid,tenant,r['id'],token,text(d.get('label') or 'NFC card',80),None,1,now(),None));return {'id':iid,'url':base_url()+'/n/'+token,'employee':json.loads(r['data']).get('name')}
        if method=='POST' and path=='/api/nfc/confirm':
            r=scoped(c,'nfc_cards',tenant,d.get('id'))
            if not r['active']:raise Problem('NFC assignment inactive.')
            tag=text(d.get('tag_uid'),100)
            if not re.fullmatch(r'[a-fA-F0-9:]+',tag):raise Problem('Invalid tag ID.')
            digest=hashlib.sha256((company(c,tenant)['analytics_salt']+tag.lower()).encode()).hexdigest();c.execute('UPDATE nfc_cards SET tag_hash=?,written=? WHERE tenant_id=? AND id=?',(digest,now(),tenant,r['id']));audit(c,actor,'nfc.device_reported_write',r['id']);return {'ok':True}
        if method=='PUT' and path.startswith('/api/nfc/cards/'):
            r=scoped(c,'nfc_cards',tenant,path.rsplit('/',1)[-1]);cid=r['card_id']
            if d.get('card_id'):cid=scoped(c,'cards',tenant,d['card_id'])['id']
            elif d.get('user_id'):
                card=c.execute('SELECT * FROM cards WHERE tenant_id=? AND user_id=? AND is_default=1',(tenant,d['user_id'])).fetchone()
                if not card:raise Problem('Member not found in this organization.',404)
                cid=card['id']
            target=scoped(c,'cards',tenant,cid);public_card(c,target['slug']);c.execute('UPDATE nfc_cards SET card_id=?,active=? WHERE tenant_id=? AND id=?',(cid,int(bool(d.get('active',r['active']))),tenant,r['id']));return {'ok':True}
        raise Problem('Route not found.',404)
    def record_activity(self,c,card,kind,source='link',label=''):
        if self.headers.get('DNT')=='1' or self.headers.get('Sec-GPC')=='1':return None
        try:
            self.actor(c)
            if kind!='share':return None
        except Problem:pass
        tenant=card['tenant_id'];cookies=SimpleCookie(self.headers.get('Cookie',''));visitor=cookies['hello_visitor'].value if 'hello_visitor' in cookies else secrets.token_urlsafe(24)
        if len(visitor)>100:visitor=secrets.token_urlsafe(24)
        digest=hashlib.sha256((company(c,tenant)['analytics_salt']+visitor).encode()).hexdigest();source=source if source in ('link','qr','nfc','staff') else 'link'
        if not c.execute('SELECT 1 FROM activity WHERE tenant_id=? AND card_id=? AND kind=? AND source=? AND visitor=? AND label=? AND created>?',(tenant,card['id'],kind,source,digest,label,now()-10)).fetchone():c.execute('INSERT INTO activity(tenant_id,user_id,card_id,kind,source,visitor,label,created) VALUES(?,?,?,?,?,?,?,?)',(tenant,card['user_id'],card['id'],kind,source,digest,label,now()))
        c.execute('DELETE FROM activity WHERE tenant_id=? AND created<?',(tenant,now()-365*86400));return 'hello_visitor='+visitor+'; HttpOnly; SameSite=Lax; Path=/; Max-Age=31536000'+('; Secure' if base_url().startswith('https://') else '')
    def csv_response(self,rows,fields,filename):
        stream=io.StringIO();writer=csv.DictWriter(stream,fieldnames=fields,extrasaction='ignore');writer.writeheader()
        for row in rows:
            row={k:("'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v) for k,v in row.items()};writer.writerow(row)
        self.send(200,stream.getvalue(),'text/csv; charset=utf-8',{'Content-Disposition':'attachment; filename="'+filename+'"'})
    def leads(self,c,actor,export):
        tenant=actor['tenant_id']
        if export:billing.require(sys.modules[__name__],c,tenant,'lead_export')
        sql='SELECT * FROM contact_exchanges WHERE tenant_id=?';args=[tenant]
        if actor['role']=='member':sql+=' AND user_id=?';args.append(actor['id'])
        rows=[{k:r[k] for k in ('id','user_id','card_id','consent_version','consented_at','source','created')}|json.loads(r['data']) for r in c.execute(sql+' ORDER BY created DESC',args)]
        if export:self.csv_response(rows,['name','email','phone','company','message','user_id','card_id','consent_version','consented_at','created'],'huddle-leads.csv');return
        return {'leads':rows}
    def analytics(self,c,actor,export):
        tenant=actor['tenant_id']
        if export:billing.require(sys.modules[__name__],c,tenant,'lead_export')
        query=self.query()
        try:days=max(1,min(365,int(query.get('days',['30'])[0])))
        except ValueError:raise Problem('Invalid date range.')
        division=query.get('division',[''])[0]
        if division:scoped(c,'divisions',tenant,division)
        sql='SELECT * FROM users WHERE tenant_id=?';args=[tenant]
        if actor['role']=='member':sql+=' AND id=?';args.append(actor['id'])
        if division:sql+=' AND division_id=?';args.append(division)
        users=c.execute(sql,args).fetchall();rows=c.execute('SELECT * FROM activity WHERE tenant_id=? AND created>=?',(tenant,now()-days*86400)).fetchall();table=[];trend={}
        allowed={u['id'] for u in users};rows=[r for r in rows if r['user_id'] in allowed]
        for u in users:
            events=[r for r in rows if r['user_id']==u['id']];views=[r for r in events if r['kind']=='view'];p=user_dict(c,u)
            table.append({'user_id':u['id'],'name':p['name'],'division_id':u['division_id'],'active':bool(u['active']),'views':len(views),'visitors':len({r['visitor'] for r in views}),'nfc_opens':sum(r['source']=='nfc' for r in views),'qr_opens':sum(r['source']=='qr' for r in views),'downloads':sum(r['kind']=='contact_download' for r in events),'link_clicks':sum(r['kind']=='link_click' for r in events),'shares':sum(r['kind']=='share' for r in events)})
        for r in rows:
            date=time.strftime('%Y-%m-%d',time.gmtime(r['created']));trend[date]=trend.get(date,0)+int(r['kind']=='view')
        if export:self.csv_response(table,['name','division_id','active','views','visitors','nfc_opens','qr_opens','downloads','link_clicks','shares'],'huddle-activity.csv');return
        return {'days':days,'staff':table,'totals':{k:sum(r[k] for r in table) for k in ('views','visitors','nfc_opens','qr_opens','downloads','link_clicks','shares')},'trend':[{'date':d,'views':n} for d,n in sorted(trend.items())],'notes':['Engagement indicators do not prove physical taps, contact imports, meetings or sales.','Visitors are pseudonymous browsers per member. DNT/GPC and authenticated previews are excluded.']}
    def qr_png(self,svg):
        from PIL import Image,ImageDraw,ImageFont
        import xml.etree.ElementTree as ET
        root=ET.fromstring(svg);_,_,width,height=map(float,root.attrib['viewBox'].split());scale=14
        image=Image.new('RGB',(int(width*scale),int(height*scale)),'white');draw=ImageDraw.Draw(image)
        def render(node,fill='#164e40'):
            fill=node.attrib.get('fill',fill);tag=node.tag.rsplit('}',1)[-1]
            if tag=='rect':
                x=float(node.attrib.get('x',0));y=float(node.attrib.get('y',0));w=node.attrib.get('width','0');h=node.attrib.get('height','0');w=width if w=='100%' else float(w);h=height if h=='100%' else float(h)
                draw.rounded_rectangle((round(x*scale),round(y*scale),round((x+w)*scale)-1,round((y+h)*scale)-1),radius=float(node.attrib.get('rx',0))*scale,fill=fill)
            elif tag=='image':
                raw=base64.b64decode(node.attrib['href'].split(',',1)[1]);logo=Image.open(io.BytesIO(raw)).convert('RGBA');logo.thumbnail((round(float(node.attrib['width'])*scale),round(float(node.attrib['height'])*scale)))
                image.paste(logo,(round(float(node.attrib['x'])*scale),round(float(node.attrib['y'])*scale)),logo)
            elif tag=='text':draw.text((width*scale/2,float(node.attrib['y'])*scale),node.text or '',fill=fill,font=ImageFont.load_default(size=24),anchor='ms')
            for child in node:render(child,fill)
        render(root);stream=io.BytesIO();image.save(stream,format='PNG');return stream.getvalue()
    def qr_svg(self,c,tenant,target,config):
        from reportlab.graphics.barcode import qrencoder
        q=qrencoder.QRCode(None,qrencoder.QRErrorCorrectLevel.H);q.addData(target);q.make();n=q.getModuleCount();fg=config.get('foreground','#164e40');bg=config.get('background','#ffffff');frame=config.get('frame','');height=n+8+(8 if frame else 0)
        eyes=[(0,0),(n-7,0),(0,n-7)];rounded=config.get('eye_style')=='rounded'
        inside=lambda x,y:any(ex<=x<ex+7 and ey<=y<ey+7 for ex,ey in eyes)
        cells=''.join(f'<rect x="{x+4}" y="{y+4}" width="1" height="1"/>' for y in range(n) for x in range(n) if q.modules[y][x] and not (rounded and inside(x,y)))
        if rounded:
            for ex,ey in eyes:
                for inset,size,fill,rx in [(0,7,fg,.7),(1,5,bg,.4),(2,3,fg,.3)]:cells+=f'<rect x="{ex+4+inset}" y="{ey+4+inset}" width="{size}" height="{size}" fill="{fill}" rx="{rx}"/>'
        logo=''
        if config.get('logo'):
            key=company(c,tenant).get('logo');a=c.execute("SELECT storage_key FROM assets WHERE tenant_id=? AND storage_key=? AND visibility='public'",(tenant,key)).fetchone()
            if a:
                raw=object_storage.read(sys.modules[__name__],a['storage_key']);size=n*.16;left=(n+8-size)/2;logo=f'<rect x="{left-1}" y="{left-1}" width="{size+2}" height="{size+2}" fill="{bg}"/><image x="{left}" y="{left}" width="{size}" height="{size}" href="data:image/png;base64,{base64.b64encode(raw).decode()}"/>'
        label=f'<text x="{(n+8)/2}" y="{n+13}" text-anchor="middle" font-family="sans-serif" font-size="2.8" fill="{fg}">{html.escape(frame)}</text>' if frame else ''
        return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {n+8} {height}" shape-rendering="crispEdges"><rect width="100%" height="100%" fill="{bg}"/><g fill="{fg}">{cells}</g>{logo}{label}</svg>'

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--demo',action='store_true');parser.add_argument('--port',type=int,default=8765);parser.add_argument('--host',default='127.0.0.1');args=parser.parse_args();DEMO=args.demo;initialise(args.demo)
    print(f'HUDDLE ready on {args.host}:{args.port}',flush=True);ThreadingHTTPServer((args.host,args.port),Handler).serve_forever()
